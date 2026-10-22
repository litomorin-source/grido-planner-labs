from datetime import datetime
from zoneinfo import ZoneInfo
import csv
import hashlib
import re
from uuid import uuid4
from io import BytesIO
from openpyxl import load_workbook
import tempfile
from pathlib import Path

import pandas as pd
import streamlit as st

from core.config import get_github_config
from core.github_sync import GitHubSync
from core.validators import validar_maestro, validar_carrito
from core.version import APP_VERSION
from core.maestro_diff import comparar_maestros

from motor import (
    procesar_archivos,
    procesar_costo_stock,
    validar_stock,
    validar_sabores,
    validar_data,
)


ADMIN_PIN = "2468"

MAESTRO_GITHUB_PATH = "data/Maestro_Productos_Grido.xlsx"
CARRITO_GITHUB_PATH = "data/Modelo_de_Carrito.xlsx"
HISTORIAL_STOCK_PATH = "data/Historial_Valorizaciones_Stock.csv"
COLUMNAS_HISTORIAL_STOCK = [
    "id", "fecha_hora", "fecha_stock", "nombre_archivo", "valor_stock", "archivo_stock_sha256"
]


def fecha_desde_nombre_stock(nombre):
    """SmartFran: Stock_YYYYMMDD_HHMMSS.csv. Hora local de Argentina."""
    nombre = Path(str(nombre)).name
    encontrado = re.fullmatch(r"Stock_(\d{8})_(\d{6})(?:\s*\(\d+\))?\.csv", nombre, re.IGNORECASE)
    if not encontrado:
        return None
    try:
        dt = datetime.strptime(''.join(encontrado.groups()[:2]), "%Y%m%d%H%M%S")
        return dt.replace(tzinfo=ZoneInfo("America/Argentina/Buenos_Aires"))
    except ValueError:
        return None


def leer_historial_stock(github):
    """El CSV vive en GitHub, no en disco efímero de Streamlit."""
    archivo = github.download_file_bytes(HISTORIAL_STOCK_PATH)
    if not archivo.get("ok"):
        info = github.get_file_info(HISTORIAL_STOCK_PATH)
        if info.get("ok") and not info.get("exists"):
            return [], None
        return [], archivo.get("message", "No se pudo leer el historial")
    try:
        contenido = archivo["content_bytes"].decode("utf-8-sig")
        filas = list(csv.DictReader(contenido.splitlines()))
        return filas, None
    except Exception as exc:
        return [], f"No se pudo interpretar el historial: {exc}"


def escribir_historial_stock(github, filas, descripcion):
    import io
    texto = io.StringIO(newline="")
    writer = csv.DictWriter(texto, fieldnames=COLUMNAS_HISTORIAL_STOCK, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(filas)
    respuesta = github.upload_bytes_file(
        path=HISTORIAL_STOCK_PATH,
        content_bytes=texto.getvalue().encode("utf-8-sig"),
        commit_message=descripcion,
    )
    if not respuesta.get("ok"):
        return False, respuesta.get("message", "Error al guardar historial")
    return True, None


def guardar_valorizacion_stock(github, valor_stock, stock_bytes, nombre_archivo):
    fecha_stock = fecha_desde_nombre_stock(nombre_archivo)
    if fecha_stock is None:
        return False, ("El nombre del archivo no tiene el formato esperado "
                       "Stock_AAAAMMDD_HHMMSS.csv. No guardé una fecha incorrecta.")
    filas, error = leer_historial_stock(github)
    if error:
        return False, error
    instante = datetime.now(ZoneInfo("America/Argentina/Buenos_Aires"))
    filas.append({
        "id": str(uuid4()),
        "fecha_hora": instante.isoformat(timespec="seconds"),
        "fecha_stock": fecha_stock.isoformat(timespec="seconds"),
        "nombre_archivo": str(nombre_archivo),
        "valor_stock": str(round(float(valor_stock), 2)),
        "archivo_stock_sha256": hashlib.sha256(stock_bytes).hexdigest(),
    })
    ok, msg = escribir_historial_stock(
        github, filas, f"Registrar valorización de stock {fecha_stock.strftime('%Y-%m-%d %H:%M')}"
    )
    if not ok:
        return False, msg
    return True, fecha_stock.strftime("%d/%m/%Y %H:%M")


def mostrar_historial_stock(github):
    filas, error = leer_historial_stock(github)
    with st.expander("Historial de valorizaciones del stock", expanded=False):
        if error:
            st.warning(f"No pude consultar el historial: {error}")
            return
        if not filas:
            st.caption("Todavía no se registraron valorizaciones de stock.")
            return
        df = pd.DataFrame(filas)
        for c in COLUMNAS_HISTORIAL_STOCK:
            if c not in df.columns:
                df[c] = ""
        df["valor_stock"] = pd.to_numeric(df["valor_stock"], errors="coerce")
        df["fecha_hora_parseada"] = pd.to_datetime(df["fecha_hora"], errors="coerce", utc=True).dt.tz_convert("America/Argentina/Buenos_Aires")
        df["fecha_stock_parseada"] = pd.to_datetime(df["fecha_stock"], errors="coerce", utc=True).dt.tz_convert("America/Argentina/Buenos_Aires")
        df = df.dropna(subset=["valor_stock"]).copy()
        df["fecha_orden"] = df["fecha_stock_parseada"].fillna(df["fecha_hora_parseada"])
        df = df.sort_values("fecha_orden", ascending=False, na_position="last")
        if df.empty:
            st.warning("El historial no tiene registros válidos.")
            return
        def fecha_visible(v):
            return v.strftime("%d/%m/%Y %H:%M") if pd.notna(v) else "—"
        mostrar = pd.DataFrame({
            "Fecha del stock": df["fecha_stock_parseada"].apply(fecha_visible),
            "Registrado el": df["fecha_hora_parseada"].apply(fecha_visible),
            "Archivo": df["nombre_archivo"].fillna("").replace("", "Archivo anterior sin nombre"),
            "Valorización del stock": df["valor_stock"].apply(formato_moneda_ar),
        })
        st.dataframe(mostrar, hide_index=True, use_container_width=True)
        st.download_button("DESCARGAR HISTORIAL CSV", data=df[COLUMNAS_HISTORIAL_STOCK].to_csv(index=False).encode("utf-8-sig"), file_name="Historial_Valorizaciones_Stock.csv", mime="text/csv", key="descargar_historial_val_stock")
        st.markdown("**Eliminar un registro**")
        opciones = {}
        for _, fila in df.iterrows():
            identificador = str(fila.get("id", ""))
            if not identificador:
                continue
            nombre = str(fila.get("nombre_archivo", "")).strip() or "Registro anterior"
            fecha = fecha_visible(fila["fecha_stock_parseada"] if pd.notna(fila["fecha_stock_parseada"]) else fila["fecha_hora_parseada"])
            etiqueta = f"{fecha} · {formato_moneda_ar(fila['valor_stock'])} · {nombre} · {identificador[:8]}"
            opciones[etiqueta] = identificador
        if opciones:
            eleccion = st.selectbox("Registro a borrar", list(opciones.keys()), key="historico_stock_elegir_borrar")
            confirmar = st.checkbox("Confirmo que quiero borrar este registro del historial", key="historico_stock_confirmar_borrar")
            if st.button("ELIMINAR REGISTRO", disabled=not confirmar, key="historico_stock_borrar"):
                elegido = opciones[eleccion]
                nuevas = [fila for fila in filas if str(fila.get("id", "")) != elegido]
                if len(nuevas) == len(filas):
                    st.error("No encontré ese registro para borrar.")
                else:
                    ok, msg = escribir_historial_stock(github, nuevas, "Eliminar registro de valorización stock")
                    if ok:
                        st.success("Registro eliminado del historial.")
                        st.rerun()
                    else:
                        st.error(f"No se pudo borrar: {msg}")


def formato_moneda_ar(valor):
    try:
        valor = int(float(valor))
    except Exception:
        valor = 0
    signo = "-" if valor < 0 else ""
    entero_txt = f"{abs(valor):,}".replace(",", ".")
    return f"$ {signo}{entero_txt}"



def _codigo_clave(valor):
    if pd.isna(valor):
        return ""
    txt = str(valor).strip()
    if txt.endswith(".0"):
        txt = txt[:-2]
    return txt


def _limpiar_categoria_carrito(valor):
    txt = str(valor or "").strip()
    if "-" in txt:
        izquierda, derecha = txt.split("-", 1)
        if izquierda.strip().upper().startswith("H"):
            txt = derecha.strip()
    return txt


def _composicion_stock_desde_carrito(detalle_stock, carrito_path):
    """
    Cruza el stock valorizado con la columna A del Carrito (GrupoProducto).

    Devuelve:
    - sabores: categorías de sabores separadas (Comunes, Premium, Al Agua, etc.)
    - general: todas las categorías con stock > 0, agrupando todos los SABORES en una sola porción.
    """
    if detalle_stock is None or detalle_stock.empty:
        return pd.DataFrame(), pd.DataFrame()

    carrito = pd.read_excel(carrito_path)
    if carrito.shape[1] < 2:
        return pd.DataFrame(), pd.DataFrame()

    mapa = carrito.iloc[:, [0, 1]].copy()
    mapa.columns = ["Categoria Carrito", "Codigo Carrito"]
    mapa["Codigo Key"] = mapa["Codigo Carrito"].apply(_codigo_clave)
    mapa["Categoria Carrito"] = mapa["Categoria Carrito"].apply(_limpiar_categoria_carrito)
    mapa = mapa[mapa["Codigo Key"] != ""].drop_duplicates("Codigo Key")

    detalle = detalle_stock.copy()
    detalle["Codigo Key"] = detalle["Código Compra"].apply(_codigo_clave)
    detalle["Stock"] = pd.to_numeric(detalle["Stock"], errors="coerce").fillna(0)

    detalle = detalle.merge(
        mapa[["Codigo Key", "Categoria Carrito"]],
        on="Codigo Key",
        how="left",
    )

    detalle = detalle[
        (detalle["Stock"] > 0)
        & detalle["Categoria Carrito"].notna()
        & (detalle["Categoria Carrito"].astype(str).str.strip() != "")
    ].copy()

    if detalle.empty:
        return pd.DataFrame(), pd.DataFrame()

    es_sabor = detalle["Categoria Carrito"].str.contains("SABORES", case=False, na=False)

    sabores = (
        detalle[es_sabor]
        .groupby("Categoria Carrito", as_index=False)["Stock"]
        .sum()
        .rename(columns={"Categoria Carrito": "Categoría"})
        .sort_values("Stock", ascending=False)
    )

    general = detalle.copy()
    general["Categoría"] = general["Categoria Carrito"].where(~es_sabor, "SABORES")
    general = (
        general.groupby("Categoría", as_index=False)["Stock"]
        .sum()
        .sort_values("Stock", ascending=False)
    )

    for df in (sabores, general):
        if not df.empty:
            total = float(df["Stock"].sum())
            df["%"] = (df["Stock"] / total * 100).round(1) if total else 0

    return sabores, general


def _grafico_torta_stock(df, key):
    if df is None or df.empty:
        st.info("No hay datos suficientes para armar este gráfico.")
        return

    chart_df = df.copy()
    chart_df["Stock"] = pd.to_numeric(chart_df["Stock"], errors="coerce").fillna(0).round(2)

    spec = {
        "mark": {"type": "arc", "innerRadius": 35},
        "encoding": {
            "theta": {"field": "Stock", "type": "quantitative"},
            "color": {
                "field": "Categoría",
                "type": "nominal",
                "legend": {"title": None},
            },
            "tooltip": [
                {"field": "Categoría", "type": "nominal", "title": "Categoría"},
                {"field": "Stock", "type": "quantitative", "title": "Stock", "format": ",.2f"},
                {"field": "%", "type": "quantitative", "title": "%", "format": ".1f"},
            ],
        },
        "view": {"stroke": None},
    }

    st.vega_lite_chart(
        chart_df,
        spec,
        use_container_width=True,
        key=key,
    )


def generar_carrito_importacion(pedido, carrito_path):
    """Conserva todas las columnas del Carrito y exporta solo filas con pedido.

    Código se guarda como texto y Cantidad como número entero (packs).
    """
    import math

    libro = load_workbook(carrito_path)
    hoja = libro.active
    columnas = {str(c.value).strip(): c.column for c in hoja[1] if c.value is not None}
    for necesaria in ("Codigo", "Cantidad"):
        if necesaria not in columnas:
            raise ValueError(f"La plantilla del Carrito no tiene la columna {necesaria}.")

    pedidos = {}
    for _, fila in pedido.iterrows():
        codigo = _codigo_clave(fila.get("Código Compra"))
        cantidad = pd.to_numeric(fila.get("Packs a Comprar"), errors="coerce")
        if not codigo or pd.isna(cantidad) or cantidad <= 0:
            continue
        if not math.isclose(float(cantidad), round(float(cantidad)), abs_tol=1e-9):
            raise ValueError(f"La cantidad del código {codigo} no es entera: {cantidad}")
        pedidos[codigo] = pedidos.get(codigo, 0) + int(round(float(cantidad)))

    encontrados = set()
    filas_sin_pedido = []
    for nro_fila in range(2, hoja.max_row + 1):
        codigo_cell = hoja.cell(row=nro_fila, column=columnas["Codigo"])
        cantidad_cell = hoja.cell(row=nro_fila, column=columnas["Cantidad"])
        codigo = _codigo_clave(codigo_cell.value)
        if codigo not in pedidos:
            filas_sin_pedido.append(nro_fila)
            continue

        # El sistema exige código como texto y cantidad como número.
        codigo_cell.value = codigo
        codigo_cell.number_format = "@"
        cantidad_cell.value = pedidos[codigo]
        cantidad_cell.number_format = "0"
        encontrados.add(codigo)

    faltantes = sorted(set(pedidos) - encontrados)
    if faltantes:
        raise ValueError(
            "Hay códigos del pedido que no aparecen en la plantilla del Carrito: "
            + ", ".join(faltantes[:20])
            + (" ..." if len(faltantes) > 20 else "")
            + ". Actualizá el Carrito antes de generar el archivo de importación."
        )

    # Se eliminan solamente filas sin cantidad; se conservan todas las columnas
    # y los encabezados y estilos de la plantilla del sistema de compras.
    for nro_fila in reversed(filas_sin_pedido):
        hoja.delete_rows(nro_fila, 1)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue(), len(encontrados), sum(pedidos.values())

def status_line(ok: bool, label: str, detail: str = ""):
    if ok:
        st.success(f"✅ {label}" + (f" — {detail}" if detail else ""))
    else:
        st.error(f"❌ {label}" + (f" — {detail}" if detail else ""))


def short_sha(sha):
    return sha[:7] if sha else "-"


def format_size(size_bytes):
    if size_bytes is None:
        return "-"
    if size_bytes < 1024:
        return f"{size_bytes} B"
    if size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    return f"{size_bytes / (1024 * 1024):.2f} MB"


def normalize(value):
    import unicodedata
    import re
    s = str(value or "").strip().lower()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def detectar_grupos_powerbi(file_path):
    try:
        df = pd.read_excel(file_path)
        cols_norm = {normalize(c): c for c in df.columns}

        # Formato nuevo P1/P2
        col_categoria = next((orig for norm, orig in cols_norm.items() if norm == "categoria"), None)
        col_subcategoria = next((orig for norm, orig in cols_norm.items() if norm == "subcategoria"), None)
        col_grupo = next((orig for norm, orig in cols_norm.items() if norm == "grupo"), None)
        col_producto = next((orig for norm, orig in cols_norm.items() if norm == "producto"), None)
        col_p1 = next((orig for norm, orig in cols_norm.items() if norm == "cant acum p1" or "cant acum p1" in norm), None)

        if all([col_categoria, col_subcategoria, col_grupo, col_producto, col_p1]):
            df = df[[col_categoria, col_subcategoria, col_grupo, col_producto]].copy()
            df.columns = ["Categoria", "SubCategoria", "Grupo", "Producto"]

        else:
            # Formato anterior
            if df.shape[1] != 16:
                return [], "El archivo no coincide con el formato anterior ni con el nuevo formato P1/P2."

            df.columns = [
                "Categoria", "SubCategoria", "Grupo", "Producto",
                "PL_Cant", "PL_Fact", "PL_Kilos", "PL_Porc",
                "Promo_Cant", "Promo_Fact", "Promo_Kilos", "Promo_Porc",
                "Total_Cantidad", "Total_Fact", "Total_Kilos", "Total_Porc"
            ]
            df = df.iloc[1:].copy()

        for c in ["Categoria", "SubCategoria", "Grupo"]:
            df[c] = df[c].ffill()

        mask = (
            ((df["Categoria"] == "Heladería") & (df["SubCategoria"] == "Impulsivos"))
            | ((df["Categoria"] == "Congelados") & (df["Grupo"].isin(["Congelados Multimarca", "Frizzio"])))
        )
        df = df[
            mask
            & df["Producto"].notna()
            & (~df["Producto"].astype(str).str.lower().eq("total"))
        ].copy()

        return sorted([str(g) for g in df["Grupo"].dropna().unique()]), ""

    except Exception as e:
        return [], f"No pude leer los grupos del Power BI. Error: {e}"


def estado_archivo(label, ok, msg, detalles=None, warning=None):
    if ok:
        st.success(f"✅ {label} OK")
        if warning:
            st.warning(warning)
        if detalles:
            with st.expander(f"Ver detalles de {label}", expanded=False):
                for item in detalles:
                    st.write(f"- {item}")
    else:
        st.error(f"❌ {label}: {msg}")
        if detalles:
            with st.expander("Ver detalle del error", expanded=True):
                st.write(detalles)


def descargar_base_github(github, tmp):
    maestro_download = github.download_file_bytes(MAESTRO_GITHUB_PATH)
    carrito_download = github.download_file_bytes(CARRITO_GITHUB_PATH)
    if not maestro_download["ok"]:
        return False, maestro_download.get("message", "Error descargando Maestro"), None, None
    if not carrito_download["ok"]:
        return False, carrito_download.get("message", "Error descargando Carrito"), None, None
    maestro_path = tmp / "Maestro_Productos_Grido.xlsx"
    carrito_path = tmp / "Modelo_de_Carrito.xlsx"
    maestro_path.write_bytes(maestro_download["content_bytes"])
    carrito_path.write_bytes(carrito_download["content_bytes"])
    ok_m, err_m, _ = validar_maestro(maestro_path)
    if not ok_m:
        return False, "El Maestro de GitHub no pasó validación: " + " | ".join(err_m), None, None
    ok_c, err_c, _ = validar_carrito(carrito_path)
    if not ok_c:
        return False, "El Carrito de GitHub no pasó validación: " + " | ".join(err_c), None, None
    return True, "Base GitHub OK", maestro_path, carrito_path


def render_centro_datos(github):
    st.header("Centro de Datos")
    if st.button("Actualizar estado", use_container_width=True):
        st.session_state["data_status"] = {
            "maestro": github.get_file_info(MAESTRO_GITHUB_PATH),
            "carrito": github.get_file_info(CARRITO_GITHUB_PATH),
            "timestamp": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
        }
    if "data_status" not in st.session_state:
        st.session_state["data_status"] = {
            "maestro": github.get_file_info(MAESTRO_GITHUB_PATH),
            "carrito": github.get_file_info(CARRITO_GITHUB_PATH),
            "timestamp": datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC"),
        }
    st.caption(f"Consulta UTC: {st.session_state['data_status']['timestamp']}")
    for nombre, key in [("Maestro", "maestro"), ("Carrito", "carrito")]:
        info = st.session_state["data_status"][key]
        st.subheader(nombre)
        if not info["ok"]:
            st.error(info["message"])
            continue
        if not info["exists"]:
            st.warning("Archivo no encontrado en GitHub.")
            continue
        col1, col2, col3 = st.columns(3)
        col1.metric("Estado", "Cargado")
        col2.metric("Tamaño", format_size(info.get("size")))
        col3.metric("SHA", short_sha(info.get("sha")))
        commit = info.get("commit") or {}
        if commit.get("ok") and commit.get("exists"):
            st.caption(f"Última actualización: {commit.get('author_date')}")
            st.code(commit.get("message") or "-")
        if info.get("html_url"):
            st.link_button(f"Ver {nombre} en GitHub", info["html_url"], use_container_width=True)
        st.markdown("---")


def subir_archivo_validado(label, uploader_label, destino, validator, filename, github):
    st.subheader(label)
    st.caption(f"Destino en GitHub: `{destino}`")
    uploaded_file = st.file_uploader(uploader_label, type=["xlsx"], key=destino)

    if not uploaded_file:
        return

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_path = Path(tmp_dir) / filename
        file_bytes = uploaded_file.getvalue()
        tmp_path.write_bytes(file_bytes)

        ok, errores, _ = validator(tmp_path)

        if not ok:
            st.error("El archivo no pasó la validación.")
            for err in errores:
                st.error(err)
            return

        st.success("Archivo validado correctamente.")

        col1, col2 = st.columns(2)
        col1.metric("Tamaño archivo", f"{len(file_bytes) / 1024:.1f} KB")
        col2.metric("Destino", destino)

        permitir_guardar = True

        if label == "Maestro":
            st.markdown("---")
            st.subheader("Control de cambios del Maestro")

            actual = github.download_file_bytes(destino)

            if not actual.get("ok"):
                st.warning("No pude descargar el Maestro vigente para comparar. Si es la primera carga, podés guardarlo igual.")
                st.caption(actual.get("message", ""))
            else:
                actual_path = Path(tmp_dir) / "Maestro_Actual_GitHub.xlsx"
                actual_path.write_bytes(actual["content_bytes"])

                try:
                    diff = comparar_maestros(actual_path, tmp_path)
                    resumen = diff["resumen"]

                    c1, c2, c3, c4, c5 = st.columns(5)
                    c1.metric("Productos actual", resumen["productos_actual"])
                    c2.metric("Productos nuevo", resumen["productos_nuevo"])
                    c3.metric("Agregados", resumen["agregados"])
                    c4.metric("Eliminados", resumen["eliminados"])
                    c5.metric("Cambios", resumen["cambios"])

                    if resumen["eliminados"] > 0:
                        st.error("Atención: el nuevo Maestro elimina productos existentes.")
                        permitir_guardar = st.checkbox(
                            "Confirmo que revisé los productos eliminados y quiero permitir el guardado.",
                            value=False,
                            key="confirmar_eliminados_maestro",
                        )

                    if resumen["cambios"] > 0:
                        st.warning("Hay productos existentes con cambios. Revisar antes de guardar.")

                    tabs = st.tabs(["Agregados", "Eliminados", "Cambios"])

                    with tabs[0]:
                        if len(diff["agregados"]) == 0:
                            st.success("No hay productos agregados.")
                        else:
                            st.dataframe(diff["agregados"], use_container_width=True)

                    with tabs[1]:
                        if len(diff["eliminados"]) == 0:
                            st.success("No hay productos eliminados.")
                        else:
                            st.dataframe(diff["eliminados"], use_container_width=True)

                    with tabs[2]:
                        if len(diff["cambios"]) == 0:
                            st.success("No hay cambios en productos existentes.")
                        else:
                            st.dataframe(diff["cambios"], use_container_width=True)

                except Exception as e:
                    st.error("No pude comparar el Maestro nuevo contra el vigente.")
                    st.exception(e)
                    permitir_guardar = False

        if st.button(
            f"Guardar {label} en GitHub",
            type="primary",
            use_container_width=True,
            key=f"btn_{destino}",
            disabled=not permitir_guardar,
        ):
            timestamp = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
            commit_message = f"Labs 2.2.0 - actualizar {label} ({timestamp})"

            with st.spinner(f"Subiendo {label} a GitHub..."):
                result = github.upload_bytes_file(path=destino, content_bytes=file_bytes, commit_message=commit_message)

            if result["ok"]:
                st.success(f"{label} guardado en GitHub correctamente.")
                st.write("Commit:")
                st.code(result.get("commit"))
                st.session_state.pop("data_status", None)
                if result.get("html_url"):
                    st.link_button("Ver commit en GitHub", result["html_url"], use_container_width=True)
            else:
                st.error(result["message"])
                st.code(result.get("details", ""))

def _ventas_col(pedido):
    for col in pedido.columns:
        if str(col).startswith("Ventas "):
            return col
    return None


def render_diagnostico_pedido(result):
    pedido = result["pedido"].copy()
    sin_clasificar = result.get("sin_clasificar")
    posibles_faltantes = result.get("posibles_faltantes")
    stock_negativo = result.get("stock_negativo")
    ventas_col = _ventas_col(pedido)

    st.markdown("---")
    st.header("Diagnóstico previo del pedido")

    total_packs = int(pd.to_numeric(pedido.get("Packs a Comprar", 0), errors="coerce").fillna(0).sum())
    valor_stock_total = result.get("valor_stock_total", 0)
    valor_pedido_total = result.get("valor_pedido_total", 0)

    ventas_sin_stock = pedido.iloc[0:0].copy()
    stock_sin_ventas = pedido.iloc[0:0].copy()

    if ventas_col:
        ventas_num = pd.to_numeric(pedido[ventas_col], errors="coerce").fillna(0)
        stock_num = pd.to_numeric(pedido["Stock"], errors="coerce").fillna(0)
        ventas_sin_stock = pedido[(ventas_num > 0) & (stock_num <= 0)].copy()
        stock_sin_ventas = pedido[(stock_num > 0) & (ventas_num == 0)].copy()

    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Productos detectados", len(pedido))
    m2.metric("Packs sugeridos", total_packs)
    m3.metric("Valor pedido", formato_moneda_ar(valor_pedido_total))
    m4.metric("Valor stock", formato_moneda_ar(valor_stock_total))

    a1, a2, a3, a4 = st.columns(4)
    a1.metric("Stock negativo", len(stock_negativo) if stock_negativo is not None else 0)
    a2.metric("Sin clasificar", len(sin_clasificar) if sin_clasificar is not None else 0)
    a3.metric("Posibles faltantes", len(posibles_faltantes) if posibles_faltantes is not None else 0)
    a4.metric("Ventas sin stock", len(ventas_sin_stock))

    if valor_pedido_total and valor_stock_total:
        relacion = float(valor_pedido_total) / max(float(valor_stock_total), 1)
        if relacion > 1:
            st.warning(f"El pedido sugerido supera el valor del stock actual. Relación pedido/stock: {relacion:.2f}. Revisar período y parámetros.")
        elif relacion > 0.6:
            st.info(f"El pedido sugerido representa {relacion:.2f} del valor del stock actual.")

    if stock_negativo is not None and len(stock_negativo) > 0:
        st.error("Hay stock negativo. Revisar inventario o descarga antes de confirmar el pedido.")

    if sin_clasificar is not None and len(sin_clasificar) > 0:
        st.warning("Hay productos sin clasificar. Puede faltar actualizar el Maestro.")

    if posibles_faltantes is not None and len(posibles_faltantes) > 0:
        st.warning("Hay posibles faltantes: productos sin ventas y con stock bajo.")

    tabs = st.tabs(["Top reposición", "Menor cobertura", "Ventas sin stock", "Stock sin ventas", "Alertas"])

    with tabs[0]:
        top_repo = pedido.copy()
        top_repo["__packs"] = pd.to_numeric(top_repo.get("Packs a Comprar", 0), errors="coerce").fillna(0)
        top_repo = top_repo.sort_values("__packs", ascending=False).head(20)
        cols = ["Grupo", "Producto", "Stock", ventas_col, "Venta Semanal", "Semanas Stock", "Packs a Comprar", "Valor Pedido Sugerido", "Código Compra", "Producto Compra"]
        cols = [c for c in cols if c and c in top_repo.columns]
        st.dataframe(top_repo[cols], use_container_width=True)

    with tabs[1]:
        cobertura = pedido.copy()
        cobertura["__semanas"] = pd.to_numeric(cobertura.get("Semanas Stock", None), errors="coerce")
        cobertura = cobertura[cobertura["__semanas"].notna()].sort_values("__semanas", ascending=True).head(20)
        cols = ["Grupo", "Producto", "Stock", ventas_col, "Venta Semanal", "Semanas Stock", "Packs a Comprar", "Valor Pedido Sugerido"]
        cols = [c for c in cols if c and c in cobertura.columns]
        st.dataframe(cobertura[cols], use_container_width=True)

    with tabs[2]:
        if len(ventas_sin_stock) == 0:
            st.success("No se detectaron productos con ventas y stock cero/negativo.")
        else:
            cols = ["Grupo", "Producto", "Stock", ventas_col, "Venta Semanal", "Packs a Comprar", "Código Compra", "Producto Compra"]
            cols = [c for c in cols if c and c in ventas_sin_stock.columns]
            st.dataframe(ventas_sin_stock[cols], use_container_width=True)

    with tabs[3]:
        if len(stock_sin_ventas) == 0:
            st.success("No se detectaron productos con stock y ventas cero.")
        else:
            cols = ["Grupo", "Producto", "Stock", ventas_col, "Semanas Stock", "Observación", "Código Compra", "Producto Compra"]
            cols = [c for c in cols if c and c in stock_sin_ventas.columns]
            st.dataframe(stock_sin_ventas[cols].head(100), use_container_width=True)

    with tabs[4]:
        hay_alertas = False
        if stock_negativo is not None and len(stock_negativo) > 0:
            hay_alertas = True
            with st.expander("Stock negativo", expanded=True):
                st.dataframe(stock_negativo, use_container_width=True)

        if posibles_faltantes is not None and len(posibles_faltantes) > 0:
            hay_alertas = True
            with st.expander("Posibles faltantes", expanded=True):
                st.dataframe(posibles_faltantes, use_container_width=True)

        if sin_clasificar is not None and len(sin_clasificar) > 0:
            hay_alertas = True
            with st.expander("Sin clasificar", expanded=True):
                st.dataframe(sin_clasificar, use_container_width=True)

        if not hay_alertas:
            st.success("No se detectaron alertas principales.")


def _scenario_value(result, key, default=0):
    try:
        return float(result.get(key, default))
    except Exception:
        return default


def _pedido_metrics(result):
    pedido = result["pedido"].copy()
    packs = int(pd.to_numeric(pedido.get("Packs a Comprar", 0), errors="coerce").fillna(0).sum())
    productos_con_pedido = int((pd.to_numeric(pedido.get("Packs a Comprar", 0), errors="coerce").fillna(0) > 0).sum())
    valor_pedido = _scenario_value(result, "valor_pedido_total", 0)
    valor_stock = _scenario_value(result, "valor_stock_total", 0)
    return {
        "productos_en_pedido": productos_con_pedido,
        "packs_sugeridos": packs,
        "valor_pedido": valor_pedido,
        "valor_stock": valor_stock,
    }


def render_simulador_parametros(snapshot, key_prefix="sim"):
    st.markdown("---")
    st.header("Simulador de parámetros")
    st.caption("Compara escenarios usando los mismos archivos del último análisis.")

    base_params = snapshot["base_params"]

    with st.expander("Configurar escenarios", expanded=False):
        c1, c2, c3 = st.columns(3)

        with c1:
            st.subheader("Escenario A")
            a_nombre = st.text_input("Nombre A", value="Conservador", key=f"{key_prefix}_a_nombre")
            a_sem = st.number_input("Semanas objetivo A", min_value=0.5, max_value=12.0, value=max(float(base_params["Semanas objetivo"]) - 1, 0.5), step=0.5, key=f"{key_prefix}_a_sem")
            a_rep = st.number_input("Reposición A", min_value=0.0, max_value=8.0, value=float(base_params["Tiempo de reposición"]), step=0.5, key=f"{key_prefix}_a_rep")
            a_dias = st.number_input("Días A", min_value=1, max_value=60, value=int(base_params["Días analizados"]), step=1, key=f"{key_prefix}_a_dias")

        with c2:
            st.subheader("Escenario B")
            b_nombre = st.text_input("Nombre B", value="Actual", key=f"{key_prefix}_b_nombre")
            b_sem = st.number_input("Semanas objetivo B", min_value=0.5, max_value=12.0, value=float(base_params["Semanas objetivo"]), step=0.5, key=f"{key_prefix}_b_sem")
            b_rep = st.number_input("Reposición B", min_value=0.0, max_value=8.0, value=float(base_params["Tiempo de reposición"]), step=0.5, key=f"{key_prefix}_b_rep")
            b_dias = st.number_input("Días B", min_value=1, max_value=60, value=int(base_params["Días analizados"]), step=1, key=f"{key_prefix}_b_dias")

        with c3:
            st.subheader("Escenario C")
            c_nombre = st.text_input("Nombre C", value="Mayor cobertura", key=f"{key_prefix}_c_nombre")
            c_sem = st.number_input("Semanas objetivo C", min_value=0.5, max_value=12.0, value=min(float(base_params["Semanas objetivo"]) + 1, 12.0), step=0.5, key=f"{key_prefix}_c_sem")
            c_rep = st.number_input("Reposición C", min_value=0.0, max_value=8.0, value=float(base_params["Tiempo de reposición"]), step=0.5, key=f"{key_prefix}_c_rep")
            c_dias = st.number_input("Días C", min_value=1, max_value=60, value=int(base_params["Días analizados"]), step=1, key=f"{key_prefix}_c_dias")

    if not st.button("SIMULAR ESCENARIOS", use_container_width=True, key=f"{key_prefix}_button_1"):
        return

    escenarios = [
        {"nombre": a_nombre, "Semanas objetivo": a_sem, "Tiempo de reposición": a_rep, "Días analizados": a_dias},
        {"nombre": b_nombre, "Semanas objetivo": b_sem, "Tiempo de reposición": b_rep, "Días analizados": b_dias},
        {"nombre": c_nombre, "Semanas objetivo": c_sem, "Tiempo de reposición": c_rep, "Días analizados": c_dias},
    ]

    resultados = []
    pedidos_por_escenario = {}

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)

        stock_path = tmp / "stock.csv"
        sabores_path = tmp / "cajas_por_sabor.xlsx"
        data_path = tmp / "data.xlsx"
        maestro_path = tmp / "Maestro_Productos_Grido.xlsx"
        carrito_path = tmp / "Modelo_de_Carrito.xlsx"

        stock_path.write_bytes(snapshot["stock_bytes"])
        sabores_path.write_bytes(snapshot["sabores_bytes"])
        data_path.write_bytes(snapshot["data_bytes"])
        maestro_path.write_bytes(snapshot["maestro_bytes"])
        carrito_path.write_bytes(snapshot["carrito_bytes"])

        with st.spinner("Simulando escenarios..."):
            for i, esc in enumerate(escenarios, start=1):
                out = tmp / f"Pedido_Simulacion_{i}.xlsx"
                res = procesar_archivos(
                    stock_file=stock_path,
                    sabores_file=sabores_path,
                    data_file=data_path,
                    maestro_file=maestro_path,
                    output_file=out,
                    overrides={
                        "Semanas objetivo": esc["Semanas objetivo"],
                        "Tiempo de reposición": esc["Tiempo de reposición"],
                        "Días analizados": esc["Días analizados"],
                    },
                    carrito_file=carrito_path,
                )

                metrics = _pedido_metrics(res)
                resultados.append({
                    "Escenario": esc["nombre"],
                    "Semanas objetivo": esc["Semanas objetivo"],
                    "Reposición": esc["Tiempo de reposición"],
                    "Días analizados": esc["Días analizados"],
                    "Productos en pedido": metrics["productos_en_pedido"],
                    "Packs sugeridos": metrics["packs_sugeridos"],
                    "Valor pedido": int(metrics["valor_pedido"]),
                    "Valor stock": int(metrics["valor_stock"]),
                })

                tmp_pedido = res["pedido"].copy()
                tmp_pedido["Escenario"] = esc["nombre"]
                pedidos_por_escenario[esc["nombre"]] = tmp_pedido

    resumen = pd.DataFrame(resultados)

    base_nombre = b_nombre
    base_row = resumen[resumen["Escenario"] == base_nombre].iloc[0]

    resumen["Dif. valor pedido vs Actual"] = resumen["Valor pedido"] - int(base_row["Valor pedido"])
    resumen["Dif. packs vs Actual"] = resumen["Packs sugeridos"] - int(base_row["Packs sugeridos"])

    st.subheader("Comparación de escenarios")

    resumen_mostrar = resumen.copy()
    for col in ["Valor pedido", "Valor stock", "Dif. valor pedido vs Actual"]:
        resumen_mostrar[col] = resumen_mostrar[col].apply(formato_moneda_ar)

    st.dataframe(resumen_mostrar, use_container_width=True)

    col_a, col_b, col_c = st.columns(3)

    for idx, row in resumen.iterrows():
        with [col_a, col_b, col_c][idx]:
            st.metric(row["Escenario"], formato_moneda_ar(row["Valor pedido"]), formato_moneda_ar(row["Dif. valor pedido vs Actual"]))
            st.caption(f"Packs: {int(row['Packs sugeridos'])} ({int(row['Dif. packs vs Actual']):+d} vs Actual)")

    st.subheader("Productos que más cambian")

    try:
        base_pedido = pedidos_por_escenario[base_nombre][["Producto", "Packs a Comprar", "Valor Pedido Sugerido"]].copy()
        base_pedido = base_pedido.rename(columns={"Packs a Comprar": "Packs Actual", "Valor Pedido Sugerido": "Valor Actual"})

        comparaciones = []

        for esc_nombre, pedido_esc in pedidos_por_escenario.items():
            if esc_nombre == base_nombre:
                continue

            p = pedido_esc[["Producto", "Packs a Comprar", "Valor Pedido Sugerido"]].copy()
            p = p.rename(columns={"Packs a Comprar": f"Packs {esc_nombre}", "Valor Pedido Sugerido": f"Valor {esc_nombre}"})

            merged = base_pedido.merge(p, on="Producto", how="outer").fillna(0)
            merged["Escenario comparado"] = esc_nombre
            merged["Dif. packs"] = pd.to_numeric(merged[f"Packs {esc_nombre}"], errors="coerce").fillna(0) - pd.to_numeric(merged["Packs Actual"], errors="coerce").fillna(0)
            merged["Dif. valor"] = pd.to_numeric(merged[f"Valor {esc_nombre}"], errors="coerce").fillna(0) - pd.to_numeric(merged["Valor Actual"], errors="coerce").fillna(0)

            comparaciones.append(merged)

        if comparaciones:
            cambios = pd.concat(comparaciones, ignore_index=True)
            cambios["Abs dif valor"] = cambios["Dif. valor"].abs()
            cambios = cambios.sort_values("Abs dif valor", ascending=False).head(30)

            cambios_mostrar = cambios[["Escenario comparado", "Producto", "Packs Actual", "Dif. packs", "Valor Actual", "Dif. valor"]].copy()
            cambios_mostrar["Valor Actual"] = cambios_mostrar["Valor Actual"].apply(formato_moneda_ar)
            cambios_mostrar["Dif. valor"] = cambios_mostrar["Dif. valor"].apply(formato_moneda_ar)

            st.dataframe(cambios_mostrar, use_container_width=True)
    except Exception as e:
        st.warning(f"No pude calcular el detalle de cambios por producto: {e}")
def render_generar_pedido(github):
    st.header("Generar pedido")
    st.caption("Subí Stock, Sabores y Power BI. El Maestro y el Carrito vigentes se toman automáticamente.")

    if "ultimo_analisis_snapshot" in st.session_state:
        with st.expander("Simulador del último análisis", expanded=True):
            render_simulador_parametros(st.session_state["ultimo_analisis_snapshot"], key_prefix="sim_1")
    col1, col2, col3 = st.columns(3)
    with col1:
        stock_file = st.file_uploader("1. Archivo de STOCK", type=["csv"])
    with col2:
        sabores_file = st.file_uploader("2. Ventas de SABORES", type=["xlsx"])
    with col3:
        data_file = st.file_uploader("3. Ventas Power BI", type=["xlsx"])

    st.subheader("Configuración")
    cfg1, cfg2, cfg3 = st.columns(3)
    with cfg1:
        semanas_objetivo = st.number_input("Semanas objetivo", min_value=0.5, max_value=12.0, value=4.0, step=0.5)
    with cfg2:
        tiempo_reposicion = st.number_input("Tiempo de reposición", min_value=0.0, max_value=8.0, value=1.0, step=0.5)
    with cfg3:
        dias_analizados = st.number_input("Días analizados", min_value=1, max_value=60, value=14, step=1)

    st.markdown("---")
    st.subheader("Estado de archivos")
    ready = True
    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp = Path(tmp_dir)
        paths = {}
        with st.spinner("Cargando Maestro y Carrito vigentes..."):
            base_ok, base_msg, maestro_path, carrito_path = descargar_base_github(github, tmp)
        if base_ok:
            def _fecha_ultimo_cambio(path):
                try:
                    info = github.get_file_info(path)
                    commit = (info or {}).get("commit") or {}
                    fecha = commit.get("author_date")
                    if fecha:
                        dt = datetime.fromisoformat(str(fecha).replace("Z", "+00:00"))
                        return dt.strftime("%d/%m/%Y")
                except Exception:
                    pass
                return "sin fecha"

            fecha_maestro = _fecha_ultimo_cambio(MAESTRO_GITHUB_PATH)
            fecha_carrito = _fecha_ultimo_cambio(CARRITO_GITHUB_PATH)

            st.success(f"✅ Maestro OK · Último cambio: {fecha_maestro}")
            st.success(f"✅ Carrito OK · Último cambio: {fecha_carrito}")
        else:
            st.error(base_msg)
            st.stop()

        if stock_file:
            p = tmp / "stock.csv"
            p.write_bytes(stock_file.getvalue())
            ok, msg = validar_stock(p)
            estado_archivo("Stock", ok, msg)
            ready = ready and ok
            paths["stock"] = p
        else:
            ready = False
            st.info("📄 Stock: pendiente")

        if sabores_file:
            p = tmp / "cajas_por_sabor.xlsx"
            p.write_bytes(sabores_file.getvalue())
            ok, msg = validar_sabores(p)
            estado_archivo("Sabores", ok, msg)
            ready = ready and ok
            paths["sabores"] = p
        else:
            ready = False
            st.info("📄 Sabores: pendiente")

        if data_file:
            p = tmp / "data.xlsx"
            p.write_bytes(data_file.getvalue())
            ok, msg = validar_data(p)
            grupos, grupo_error = detectar_grupos_powerbi(p)
            detalles = []
            if grupos:
                detalles.append("Grupos detectados en la exportación:")
                detalles.extend(grupos)
            if grupo_error:
                detalles.append(grupo_error)
            warning = "Revisá que estén todos los grupos esperados. Si no desplegaste el '+', el pedido puede quedar incompleto."
            estado_archivo("Power BI", ok, msg, detalles=detalles if detalles else None, warning=warning if ok else None)
            ready = ready and ok
            paths["data"] = p
        else:
            ready = False
            st.info("📄 Power BI: pendiente")

        st.markdown("---")
        st.subheader("Informe costo stock")
        mostrar_historial_stock(github)
        if "stock" in paths:
            if st.button("GENERAR INFORME COSTO STOCK", use_container_width=True):
                output_stock = tmp / "Informe_Costo_Stock.xlsx"
                try:
                    result_stock = procesar_costo_stock(
                        stock_file=paths["stock"],
                        maestro_file=maestro_path,
                        output_file=output_stock,
                        carrito_file=carrito_path,
                    )
                    st.success("Informe de costo de stock generado correctamente.")
                    valor_total_stock = result_stock.get("valor_stock_total", 0)
                    guardado, resultado_guardado = guardar_valorizacion_stock(
                        github, valor_total_stock, stock_file.getvalue(), stock_file.name
                    )
                    if guardado:
                        st.success(f"Valorización guardada en el historial · {resultado_guardado}")
                    else:
                        st.error(f"El informe se generó, pero NO se guardó en el historial: {resultado_guardado}")
                    st.metric("Valor stock actual", formato_moneda_ar(valor_total_stock))
                    if result_stock.get("categoria") is not None:
                        with st.expander("Ver valorización por categoría", expanded=False):
                            df_categoria = result_stock["categoria"].copy()
                            if "Valor Stock Actual" in df_categoria.columns:
                                df_categoria["Valor Stock Actual"] = df_categoria["Valor Stock Actual"].apply(formato_moneda_ar)
                            st.dataframe(df_categoria, use_container_width=True)

                    sabores_stock, composicion_general = _composicion_stock_desde_carrito(
                        result_stock.get("detalle"),
                        carrito_path,
                    )

                    st.markdown("#### Composición del stock")
                    graf1, graf2 = st.columns(2)

                    with graf1:
                        st.markdown("**Cajas de sabores por categoría**")
                        _grafico_torta_stock(
                            sabores_stock,
                            key="stock_torta_sabores",
                        )

                    with graf2:
                        st.markdown("**Stock por categoría del Carrito**")
                        st.caption("Todos los tipos de sabores se muestran sumados como SABORES.")
                        _grafico_torta_stock(
                            composicion_general,
                            key="stock_torta_general",
                        )

                    st.download_button(
                        "DESCARGAR INFORME COSTO STOCK",
                        data=output_stock.read_bytes(),
                        file_name="Informe_Costo_Stock.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                    )
                except Exception as e:
                    st.error("Ocurrió un error al generar el informe de costo stock.")
                    st.exception(e)

        if not ready:
            st.info("Cuando los 3 archivos estén en OK, se habilita la generación del pedido.")
            st.stop()

        st.success("✅ Todo listo para generar el pedido.")
        if st.button("ANALIZAR Y GENERAR PEDIDO", type="primary", use_container_width=True):
            output = tmp / "Pedido_Final.xlsx"
            try:
                result = procesar_archivos(
                    stock_file=paths["stock"],
                    sabores_file=paths["sabores"],
                    data_file=paths["data"],
                    maestro_file=maestro_path,
                    output_file=output,
                    overrides={
                        "Semanas objetivo": semanas_objetivo,
                        "Tiempo de reposición": tiempo_reposicion,
                        "Días analizados": dias_analizados,
                    },
                    carrito_file=carrito_path,
                )
                pedido = result["pedido"]
                sin_clasificar = result["sin_clasificar"]
                posibles_faltantes = result.get("posibles_faltantes")
                stock_negativo = result.get("stock_negativo")
                st.success("Pedido analizado correctamente.")
                render_diagnostico_pedido(result)

                st.session_state["ultimo_analisis_snapshot"] = {
                    "stock_bytes": stock_file.getvalue(),
                    "sabores_bytes": sabores_file.getvalue(),
                    "data_bytes": data_file.getvalue(),
                    "maestro_bytes": maestro_path.read_bytes(),
                    "carrito_bytes": carrito_path.read_bytes(),
                    "base_params": {
                        "Semanas objetivo": semanas_objetivo,
                        "Tiempo de reposición": tiempo_reposicion,
                        "Días analizados": dias_analizados,
                    },
                }

                render_simulador_parametros(st.session_state["ultimo_analisis_snapshot"], key_prefix="sim_2")

                st.markdown("---")
                st.header("Resumen final")
                m1, m2, m3, m4, m5 = st.columns(5)
                m1.metric("Productos en pedido", len(pedido))
                m2.metric("Stock negativo", len(stock_negativo) if stock_negativo is not None else 0)
                m3.metric("Sin clasificar", len(sin_clasificar))
                m4.metric("Posibles faltantes", len(posibles_faltantes) if posibles_faltantes is not None else 0)
                m5.metric("Packs sugeridos", int(pedido["Packs a Comprar"].fillna(0).sum()))
                v1, v2 = st.columns(2)
                v1.metric("Valor stock actual", formato_moneda_ar(result.get("valor_stock_total", 0)))
                v2.metric("Valor pedido sugerido", formato_moneda_ar(result.get("valor_pedido_total", 0)))
                if result.get("valorizacion_categoria") is not None:
                    with st.expander("Ver valorización por categoría", expanded=False):
                        df_categoria = result["valorizacion_categoria"].copy()
                        if "Valor Stock Actual" in df_categoria.columns:
                            df_categoria["Valor Stock Actual"] = df_categoria["Valor Stock Actual"].apply(formato_moneda_ar)
                        st.dataframe(df_categoria, use_container_width=True)
                if result.get("valorizacion_producto") is not None:
                    with st.expander("Ver valorización por producto", expanded=False):
                        st.dataframe(result["valorizacion_producto"], use_container_width=True)
                if stock_negativo is not None and len(stock_negativo) > 0:
                    st.error("⚠️ Se detectaron productos con stock negativo.")
                if posibles_faltantes is not None and len(posibles_faltantes) > 0:
                    st.warning("Se detectaron posibles faltantes.")
                if len(sin_clasificar) > 0:
                    st.warning("Hay productos sin clasificar.")
                st.download_button(
                    "DESCARGAR PEDIDO FINAL",
                    data=output.read_bytes(),
                    file_name="Pedido_Final.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                )
                try:
                    carrito_importable, cantidad_productos, total_packs = generar_carrito_importacion(
                        pedido, carrito_path
                    )
                    st.download_button(
                        "DESCARGAR CARRITO PARA IMPORTAR",
                        data=carrito_importable,
                        file_name="Carrito_Importar.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True,
                        key="descargar_carrito_importacion",
                    )
                    st.caption(
                        f"Carrito importable: {cantidad_productos} productos · {total_packs} packs. "
                        "Solo productos con cantidad > 0; todas las columnas originales. Código como texto y cantidad numérica."
                    )
                except Exception as exc_importacion:
                    st.warning(f"No se pudo preparar el carrito para importar: {exc_importacion}")
                with st.expander("Ver vista previa del pedido", expanded=False):
                    st.dataframe(pedido.head(100), use_container_width=True)
                if stock_negativo is not None and len(stock_negativo) > 0:
                    with st.expander("Ver stock negativo", expanded=True):
                        st.dataframe(stock_negativo, use_container_width=True)
                if posibles_faltantes is not None and len(posibles_faltantes) > 0:
                    with st.expander("Ver posibles faltantes", expanded=True):
                        st.dataframe(posibles_faltantes, use_container_width=True)
                if len(sin_clasificar) > 0:
                    with st.expander("Ver productos sin clasificar", expanded=True):
                        st.dataframe(sin_clasificar, use_container_width=True)
            except Exception as e:
                st.error("Ocurrió un error al generar el pedido.")
                st.exception(e)



def render_historial_maestro(github):
    st.subheader("Historial del Maestro")
    st.caption(f"Archivo: `{MAESTRO_GITHUB_PATH}`")

    col_a, col_b = st.columns([2, 1])
    with col_a:
        cantidad = st.number_input("Cantidad de versiones a mostrar", min_value=3, max_value=30, value=10, step=1)
    with col_b:
        st.write("")
        st.write("")
        consultar = st.button("Consultar historial", use_container_width=True)

    if consultar or "historial_maestro" not in st.session_state:
        with st.spinner("Consultando historial del Maestro..."):
            st.session_state["historial_maestro"] = github.list_commits_for_path(MAESTRO_GITHUB_PATH, per_page=int(cantidad))

    historial = st.session_state.get("historial_maestro")

    if not historial:
        return

    if not historial.get("ok"):
        st.error(historial.get("message"))
        st.code(historial.get("details", ""))
        return

    commits = historial.get("commits", [])

    if not commits:
        st.warning("No encontré historial para el Maestro.")
        return

    st.info("Restaurar una versión anterior no borra historia: crea un commit nuevo con esa versión como Maestro vigente.")

    for idx, c in enumerate(commits):
        titulo = f"{idx + 1}. {c.get('author_date', '-') } — {c.get('short_sha', '-')}"
        with st.expander(titulo, expanded=(idx == 0)):
            st.write("Mensaje:")
            st.code(c.get("message") or "-")
            st.caption(f"Autor: {c.get('author_name') or '-'}")
            if c.get("html_url"):
                st.link_button("Ver commit en GitHub", c["html_url"], use_container_width=True)

            if idx == 0:
                st.success("Esta es la versión actual.")
                continue

            st.markdown("---")
            st.warning("Restaurar esta versión reemplazará el Maestro vigente por el archivo de este commit.")

            confirmar = st.checkbox(
                f"Confirmo que quiero habilitar restauración de {c.get('short_sha')}",
                key=f"confirm_restore_{c.get('sha')}",
                value=False,
            )

            if st.button(
                f"Restaurar versión {c.get('short_sha')}",
                key=f"restore_{c.get('sha')}",
                disabled=not confirmar,
                use_container_width=True,
            ):
                with tempfile.TemporaryDirectory() as tmp_dir:
                    tmp = Path(tmp_dir)

                    with st.spinner("Descargando versión histórica..."):
                        historico = github.download_file_bytes_at_ref(MAESTRO_GITHUB_PATH, c["sha"])

                    if not historico.get("ok"):
                        st.error(historico.get("message"))
                        st.code(historico.get("details", ""))
                        return

                    historico_path = tmp / "Maestro_Historico.xlsx"
                    historico_path.write_bytes(historico["content_bytes"])

                    ok, errores, _ = validar_maestro(historico_path)

                    if not ok:
                        st.error("La versión histórica no pasó validación. No se restaura.")
                        for err in errores:
                            st.error(err)
                        return

                    actual = github.download_file_bytes(MAESTRO_GITHUB_PATH)
                    if actual.get("ok"):
                        actual_path = tmp / "Maestro_Actual.xlsx"
                        actual_path.write_bytes(actual["content_bytes"])
                        try:
                            diff = comparar_maestros(actual_path, historico_path)
                            resumen = diff["resumen"]

                            st.subheader("Diferencias contra versión actual")
                            d1, d2, d3, d4, d5 = st.columns(5)
                            d1.metric("Productos actual", resumen["productos_actual"])
                            d2.metric("Productos restaurado", resumen["productos_nuevo"])
                            d3.metric("Agregados", resumen["agregados"])
                            d4.metric("Eliminados", resumen["eliminados"])
                            d5.metric("Cambios", resumen["cambios"])

                            with st.expander("Ver cambios de restauración", expanded=False):
                                tab_a, tab_b, tab_c = st.tabs(["Agregados", "Eliminados", "Cambios"])
                                with tab_a:
                                    st.dataframe(diff["agregados"], use_container_width=True)
                                with tab_b:
                                    st.dataframe(diff["eliminados"], use_container_width=True)
                                with tab_c:
                                    st.dataframe(diff["cambios"], use_container_width=True)
                        except Exception as e:
                            st.warning(f"No pude comparar antes de restaurar: {e}")

                    timestamp = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S UTC")
                    commit_message = f"Labs 2.3.0 - restaurar Maestro a {c.get('short_sha')} ({timestamp})"

                    with st.spinner("Restaurando Maestro en GitHub..."):
                        result = github.upload_bytes_file(
                            path=MAESTRO_GITHUB_PATH,
                            content_bytes=historico["content_bytes"],
                            commit_message=commit_message,
                        )

                    if result.get("ok"):
                        st.success("Maestro restaurado correctamente.")
                        st.write("Nuevo commit:")
                        st.code(result.get("commit"))
                        st.session_state.pop("data_status", None)
                        st.session_state.pop("historial_maestro", None)
                        if result.get("html_url"):
                            st.link_button("Ver nuevo commit", result["html_url"], use_container_width=True)
                    else:
                        st.error(result.get("message"))
                        st.code(result.get("details", ""))


def render_actualizar_archivos(github):
    tab_maestro, tab_carrito, tab_historial = st.tabs(["Actualizar Maestro", "Actualizar Carrito", "Historial Maestro"])

    with tab_maestro:
        subir_archivo_validado(
            label="Maestro",
            uploader_label="Subir Maestro_Productos_Grido.xlsx",
            destino=MAESTRO_GITHUB_PATH,
            validator=validar_maestro,
            filename="Maestro_Productos_Grido.xlsx",
            github=github,
        )

    with tab_carrito:
        subir_archivo_validado(
            label="Carrito",
            uploader_label="Subir Modelo_de_Carrito.xlsx",
            destino=CARRITO_GITHUB_PATH,
            validator=validar_carrito,
            filename="Modelo_de_Carrito.xlsx",
            github=github,
        )

    with tab_historial:
        render_historial_maestro(github)

def render_home():
    st.title("🍦 GridoPlanner Labs")
    st.caption(APP_VERSION)

    modo = st.sidebar.radio("Modo", ["Usuario", "Administrador"])
    st.sidebar.markdown("---")
    st.sidebar.info(f"Versión {APP_VERSION}")

    config = get_github_config()

    if not config["ok"]:
        st.error("Faltan Secrets de GitHub.")
        for item in config["missing"]:
            st.code(item)
        st.stop()

    github = GitHubSync(config["token"], config["repo"], config["branch"])

    if modo == "Usuario":
        render_generar_pedido(github)
        return

    st.header("Administrador")
    pin = st.text_input("PIN de administrador", type="password")

    if pin != ADMIN_PIN:
        st.warning("Ingresá el PIN correcto para administrar archivos.")
        st.stop()

    st.success("PIN correcto.")

    with st.expander("Estado GitHub", expanded=False):
        st.write("Repositorio:")
        st.code(config["repo"])
        st.write("Rama:")
        st.code(config["branch"])

        if st.button("Probar conexión GitHub", use_container_width=True):
            with st.spinner("Conectando..."):
                result = github.verify()

            repo = result.get("repo")
            branch = result.get("branch")

            status_line(repo["ok"], "Repositorio accesible", repo.get("message", ""))
            if branch:
                status_line(branch["ok"], "Rama encontrada", branch.get("message", ""))
                if branch.get("ok"):
                    st.write("Último commit:")
                    st.code(branch.get("commit_sha"))

    tab1, tab2 = st.tabs(["Centro de Datos", "Actualizar archivos"])

    with tab1:
        render_centro_datos(github)

    with tab2:
        render_actualizar_archivos(github)
